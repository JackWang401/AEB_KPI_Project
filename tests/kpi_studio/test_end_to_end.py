from pathlib import Path

import pandas as pd
import pytest

from src.kpi_studio.app.run_service import RunService
from src.kpi_studio.io import profile_store as ps

PROFILE = ps.builtin_profile_dir() / "aeb_basic.kpi.yaml"


@pytest.mark.parametrize("workers", [1, 2])
def test_run_aeb_profile(mf4_dir, tmp_path, workers):
    folder, info = mf4_dir
    files = [str(folder / "a.mf4"), str(folder / "b.mf4")]
    svc = RunService(PROFILE, files, tmp_path / "out", workers=workers)
    report = svc.preflight()
    assert report.ok, report.text()
    res = svc.run()
    assert res.n_ok == 2
    summary = pd.read_excel(res.excel, sheet_name="Summary")
    assert len(summary) == 3  # two activations in a, one in b
    assert summary["vehSpd [km/h]"].round(1).tolist() == [40.0, 40.0, 60.0]  # b is logged in km/h, converted
    assert summary["Start [s]"].tolist() == pytest.approx([5.0, 13.0, 4.0], abs=0.011)
    dead = summary["aebDeadTime [s]"]
    assert ((dead > 0) & (dead < 0.3)).all()
    sheets = pd.read_excel(res.excel, sheet_name=None)
    assert {"Summary", "Details", "Moments", "Files", "RunInfo"} <= set(sheets)
    files_sheet = sheets["Files"]
    assert (files_sheet["status"] == "ok").all()
    assert res.dashboard and len(res.plots) == 3 + 3
    assert (res.output_dir / "profile.resolved.kpi.yaml").exists()
    from openpyxl import load_workbook

    ws = load_workbook(res.excel)["Summary"]
    floats = [c for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, float)]
    assert floats and all(c.number_format == "0.000" for c in floats)  # shown with 3 decimals
    wb = load_workbook(res.excel)
    for sheet in wb.worksheets:  # headers centered, every data cell left aligned
        assert all(c.alignment.horizontal == "center" for c in sheet[1])
        assert all(c.alignment.horizontal == "left" for row in sheet.iter_rows(min_row=2) for c in row)


def test_bad_file_is_isolated(mf4_dir, tmp_path):
    folder, _ = mf4_dir
    broken = folder / "broken.mf4"
    broken.write_bytes(b"not an mf4")
    svc = RunService(PROFILE, [str(folder / "a.mf4"), str(broken)], tmp_path / "out")
    report = svc.preflight()
    assert any("broken.mf4" in e for e in report.errors)
    res = svc.run()
    statuses = {o.file.split("/")[-1]: o.status for o in res.outcomes}
    assert statuses == {"a.mf4": "ok", "broken.mf4": "error"}


def test_cli_validate_and_run(mf4_dir, tmp_path, capsys):
    from src.kpi_studio.cli import main

    folder, _ = mf4_dir
    assert main(["validate", str(PROFILE)]) == 0
    assert main(["run", str(PROFILE), str(folder), "-o", str(tmp_path / "o")]) == 0
    assert "Recordings OK: 2/2" in capsys.readouterr().out


def test_cli_offset_option(mf4_dir):
    from src.kpi_studio.cli import _offsets

    folder, _ = mf4_dir
    files = [str(folder / "a.mf4"), str(folder / "b.mf4")]
    assert _offsets(["b.mf4=0.25"], files) == {files[1]: 0.25}
    assert _offsets([f"{files[0]}=-1"], files) == {files[0]: -1.0}
    with pytest.raises(SystemExit):
        _offsets(["c.mf4=1"], files)
    with pytest.raises(SystemExit):
        _offsets(["b.mf4=soon"], files)


def test_results_next_to_input_are_replaced(mf4_dir):
    from src.kpi_studio.app.run_service import OUTPUT_MARKER, input_results_dir

    folder, _ = mf4_dir
    files = [str(folder / "a.mf4"), str(folder / "b.mf4")]
    target = input_results_dir(files)
    assert target == folder.resolve() / "kpi_results"
    first = RunService(PROFILE, files, None).run()
    assert first.output_dir == target and (target / "kpi_results.xlsx").exists()
    assert (target / OUTPUT_MARKER).exists() and first.dashboard.startswith(str(target))
    (target / "stale.txt").write_text("from the previous run")
    second = RunService(PROFILE, files, None).run()
    assert second.output_dir == target and not (target / "stale.txt").exists()  # replaced, not merged
    assert sorted(p.name for p in folder.iterdir() if p.name.startswith("kpi_results")) == ["kpi_results"]


def test_foreign_folder_named_kpi_results_is_not_deleted(mf4_dir):
    folder, _ = mf4_dir
    files = [str(folder / "a.mf4")]
    mine = folder / "kpi_results"
    mine.mkdir()
    (mine / "notes.txt").write_text("user data")
    with pytest.raises(RuntimeError, match="not written by KPI Studio"):
        RunService(PROFILE, files, None).run()
    assert (mine / "notes.txt").exists()


def test_merged_and_kpi_coloured_plots_are_written(mf4_dir, tmp_path):
    folder, _ = mf4_dir
    text = PROFILE.read_text().replace(
        "x: vehSpd, y: aebStopDistance, color_by: file }",
        "x: vehSpd, y: aebStopDistance, color_by: aebAverageAccel, merge: braking }").replace(
        "x: vehSpd, y: aebDeadTime, color_by: file }", "x: vehSpd, y: aebDeadTime, color_by: none, merge: braking }")
    prof = PROFILE.parent / f".test_merge_{tmp_path.name}.kpi.yaml"  # beside the original: its imports resolve
    prof.write_text(text)
    try:
        res = RunService(prof, [str(folder / "a.mf4"), str(folder / "b.mf4")], tmp_path / "out").run()
    finally:
        prof.unlink()
    assert res.n_ok == 2
    names = [Path(p).name for p in res.plots]
    assert sum("Stop distance" in n for n in names) == 1  # two scatter plots, one chart, the first's title
    assert not any("Dead time" in n for n in names)


def test_custom_output_folder_keeps_one_excel(mf4_dir, tmp_path):
    folder, _ = mf4_dir
    files = [str(folder / "a.mf4")]
    root = tmp_path / "mine"
    (root).mkdir()
    (root / "notes.txt").write_text("keep me")
    first = RunService(PROFILE, files, root).run()
    second = RunService(PROFILE, files, root).run()
    assert first.excel == second.excel == root / "kpi_results" / "kpi_results.xlsx" and second.excel.exists()
    assert sorted(p.name for p in root.iterdir()) == ["kpi_results", "notes.txt"]  # replaced, not added


def test_runs_collect_in_one_excel_replacing_the_same_file(mf4_dir, tmp_path):
    folder, _ = mf4_dir
    a, b = str(folder / "a.mf4"), str(folder / "b.mf4")
    root = tmp_path / "out"
    RunService(PROFILE, [a, b], root, versions={a: "v1", b: "v1"}).run()
    res = RunService(PROFILE, [a], root, versions={a: "v2"}).run()
    summary = pd.read_excel(res.excel, sheet_name="Summary")
    assert list(summary.columns[:2]) == ["File", "SW version"]
    assert sorted(summary["File"].unique()) == sorted([a, b])  # b kept from the first run
    assert len(summary) == 3 and (summary["File"] == a).sum() == 2  # a replaced, not added again
    assert set(summary[summary["File"] == a]["SW version"]) == {"v2"}
    assert set(summary[summary["File"] == b]["SW version"]) == {"v1"}
    files = pd.read_excel(res.excel, sheet_name="Files")
    assert sorted(files["file"]) == sorted([a, b])
    details = pd.read_excel(res.excel, sheet_name="Details")
    assert set(details["file"]) == {a, b} and "SW version" in details.columns
    res = RunService(PROFILE, [a], root).run()  # no version typed this time: the earlier one stays
    summary = pd.read_excel(res.excel, sheet_name="Summary")
    assert set(summary[summary["File"] == a]["SW version"]) == {"v2"}
